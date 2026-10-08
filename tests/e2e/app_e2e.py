"""End-to-end test of Block 2 in a real browser against the local stack.

Covers: purchase arrives (simulated webhook) -> sign up -> confirm email -> set up company,
warehouse and shifts -> Today -> settings -> account (change password) -> sign out/in ->
forgot/reset password -> buyer with no purchase -> supervisor invitation -> refund locks access.
Screenshots of every screen are saved for visual review.
"""
import asyncio, email, glob, html, os, re, subprocess, sys, time
from email import policy
from playwright.async_api import async_playwright, expect

BASE = 'http://localhost:5173'
API = 'http://localhost:54321'
MAIL = os.path.expanduser('~/.wap-localstack/mail')
SHOTS = sys.argv[1] if len(sys.argv) > 1 else 'screenshots'
os.makedirs(SHOTS, exist_ok=True)

BUYER = 'buyer@example.com'
STRANGER = 'nobuy@example.com'
SUPERVISOR = 'sup@example.com'
PW1, PW2, PW3 = 'FloorCount-2026!', 'NewSecret-2026!', 'ResetSecret-2026!'

passed, failed = [], []

def totp(secret, at=None):
    # RFC 6238 code, as an authenticator app computes it.
    import base64, hashlib, hmac, struct
    key = base64.b32decode(secret.upper() + '=' * (-len(secret) % 8))
    counter = int((at or time.time()) // 30)
    h = hmac.new(key, struct.pack('>Q', counter), hashlib.sha1).digest()
    o = h[-1] & 15
    return str((struct.unpack('>I', h[o:o + 4])[0] & 0x7fffffff) % 1000000).zfill(6)

OPEN_PAGES = []
PW = []
def check(name, cond, detail=''):
    (passed if cond else failed).append(name + (f' — {detail}' if detail and not cond else ''))

def psql(sql):
    r = subprocess.run(['psql', '-h', '/tmp', '-p', '54329', '-U', 'postgres', '-d', 'wap_e2e',
                        '-v', 'ON_ERROR_STOP=1', '-qAt', '-c', sql], capture_output=True, text=True)
    if r.returncode:
        raise RuntimeError(r.stderr)
    return r.stdout.strip()

def mail_link(to, kind, after, timeout=15):
    """Newest Supabase Auth email to `to` received after `after`; returns its verify link."""
    tag = to.replace('@', '_at_')
    end = time.time() + timeout
    while time.time() < end:
        files = sorted(f for f in glob.glob(f'{MAIL}/*_{tag}.eml') if float(os.path.basename(f).split('_')[0]) > after)
        for f in reversed(files):
            msg = email.message_from_string(open(f).read(), policy=policy.default)
            body = ''.join(p.get_content() for p in msg.walk() if p.get_content_type() in ('text/html', 'text/plain'))
            m = re.search(r'href="([^"]*/verify\?[^"]*type=' + kind + r'[^"]*)"', body)
            if m:
                return html.unescape(m.group(1))
        time.sleep(0.3)
    raise RuntimeError(f'no {kind} email for {to}')

async def shot(page, name):
    await page.screenshot(path=f'{SHOTS}/{name}.png', full_page=True)

async def main():
    # A purchase arrives from Hotmart before the buyer has an account.
    out = psql("select billing.apply_hotmart_event('e2e-1','PURCHASE_APPROVED', now(), 'SUB-E2E', null, "
               f"'{BUYER}', 'professional', now() + interval '30 days')")
    check('webhook creates the license', out == 'license_created', out)

    p = PW[0]
    if True:
        browser = await p.chromium.launch()
        ctx = await browser.new_context(viewport={'width': 1280, 'height': 860})
        page = await ctx.new_page()
        OPEN_PAGES.append(page)
        console_errors = []
        page.on('pageerror', lambda e: console_errors.append(str(e)))
        if os.environ.get('E2E_DEBUG'):
            page.on('console', lambda m: print('console:', m.text))
            page.on('framenavigated', lambda f: print('nav:', f.url) if f == page.main_frame else None)
            page.on('request', lambda r: print('req:', r.method, r.url[:120]) if '54321' in r.url else None)

        # ---------- sign up ----------
        await page.goto(BASE)
        await expect(page).to_have_url(re.compile(r'/login'))
        await shot(page, '01-login')
        await page.get_by_role('link', name='Create an account').click()
        await page.get_by_label('Your name').fill('Gary Tester')
        await page.get_by_label('Email').fill(BUYER)
        await page.get_by_label('Password', exact=True).fill('short')
        await shot(page, '02-signup-validation')
        check('signup blocks short passwords', await page.get_by_role('button', name='Create account').is_disabled())
        await page.get_by_label('Password', exact=True).fill(PW1)
        await page.get_by_label('Confirm password').fill(PW1)
        t0 = time.time()
        await page.get_by_role('button', name='Create account').click()
        await expect(page.get_by_role('heading', name='Check your email')).to_be_visible()
        await shot(page, '03-check-email')

        # ---------- unconfirmed login is refused ----------
        await page.goto(f'{BASE}/login')
        await page.get_by_label('Email').fill(BUYER)
        await page.get_by_label('Password').fill(PW1)
        await page.get_by_role('button', name='Sign in').click()
        await expect(page.get_by_role('alert')).to_contain_text('Confirm your email')
        check('unconfirmed email cannot sign in', True)

        # ---------- confirm email -> onboarding ----------
        await page.goto(mail_link(BUYER, 'signup', t0))
        await expect(page.get_by_role('heading', name='Set up your company')).to_be_visible(timeout=10000)
        check('confirmation link signs in and opens setup', True)
        await shot(page, '04-setup-company')
        await page.get_by_label('Company name').fill('Northgate Logistics')
        await page.get_by_label('Time zone').select_option('America/New_York')
        await page.get_by_role('button', name='Continue').click()

        await expect(page.get_by_role('heading', name='Add your warehouse')).to_be_visible()
        await shot(page, '05-setup-warehouse')
        await page.get_by_label('Warehouse name').fill('DC-1 Columbus')
        await page.get_by_role('button', name='Continue').click()

        await expect(page.get_by_role('heading', name=re.compile('Set the shifts'))).to_be_visible()
        await shot(page, '06-setup-shifts')
        # Second shift crosses midnight: the editor must say so.
        check('overnight shift is explained', await page.get_by_text('ends the next day').count() == 1)
        await page.get_by_role('button', name='Add another shift').click()
        await page.locator('#shift-2-name').fill('Weekend Shift')
        await page.locator('#shift-2-start').fill('07:00')
        await page.locator('#shift-2-end').fill('15:30')
        for d in ['Mon', 'Tue', 'Wed', 'Thu', 'Fri']:
            await page.locator('fieldset').nth(2).get_by_role('button', name=d).click()
        for d in ['Sat', 'Sun']:
            await page.locator('fieldset').nth(2).get_by_role('button', name=d).click()
        await page.get_by_role('button', name='Save shifts').click()
        await expect(page.get_by_role('heading', name='Your company is ready')).to_be_visible()
        await shot(page, '07-setup-done')
        rows = psql("select string_agg(name || ' ' || start_time || '-' || end_time || ' ' || days::text, ' | ' order by sort_order) from shifts")
        check('three shifts saved with their days', rows.count('|') == 2 and 'Weekend Shift 07:00:00-15:30:00 {6,7}' in rows, rows)

        # ---------- Today ----------
        await page.get_by_role('button', name='Go to Today').click()
        await expect(page.get_by_role('heading', name='Today', exact=True)).to_be_visible()
        await shot(page, '08-today')
        check('Today shows the warehouse', await page.get_by_role('heading', name='DC-1 Columbus').count() == 1)

        # ---------- Settings ----------
        await page.get_by_role('link', name='Settings').click()
        await expect(page.get_by_role('heading', name='Settings')).to_be_visible()
        await page.get_by_label('Company name').fill('Northgate Logistics LLC')
        await page.get_by_role('button', name='Save company').click()
        await expect(page.get_by_text('Company saved.')).to_be_visible()
        check('company renamed in database', psql('select name from organizations') == 'Northgate Logistics LLC')
        await page.locator('[id$="-0-grace"]').first.select_option('10')
        await page.get_by_role('button', name='Save shifts').click()
        await expect(page.get_by_text('Shifts saved.')).to_be_visible()
        check('shift edit saved', psql("select late_grace_minutes from shifts where name='First Shift'") == '10')
        await page.get_by_label('Warehouse name').fill('DC-2 Groveport')
        await page.get_by_role('button', name='Add warehouse').click()
        await expect(page.get_by_role('alert')).to_contain_text('plan doesn’t include another warehouse')
        check('plan limit explained when adding a 2nd warehouse', True)
        await shot(page, '09-settings')

        # ---------- Employees: import a messy Excel file ----------
        import openpyxl, datetime
        wb = openpyxl.Workbook(); sh = wb.active
        sh.append(['Northgate staffing roster'])                       # title row above the header
        sh.append(['Associate ID', 'Name', 'Shift', 'Dept', 'Hire Date'])
        sh.append([10234, 'Lopez, Maria', 'First Shift', 'Picking', datetime.date(2026, 9, 14)])
        sh.append([10235, 'James Carter', 'first shift', 'Receiving', None])
        sh.append([10236, 'Nguyen, Linh', 'Second Shift', 'Picking', None])
        sh.append([10237, 'Night, Owl', 'Night Shift', '', None])       # unknown shift
        sh.append([10234, 'Lopez, Maria', 'First Shift', 'Picking', None])  # duplicate badge
        xlsx = os.path.join(SHOTS, 'roster.xlsx'); wb.save(xlsx)

        await page.get_by_role('link', name='Employees', exact=True).click()
        await expect(page.get_by_role('heading', name=re.compile('No employees in'))).to_be_visible()
        await shot(page, '09a-employees-empty')
        await page.get_by_role('link', name='Import from Excel or CSV').first.click()
        await page.locator('input[type=file]').set_input_files(xlsx)
        await expect(page.get_by_text('3. Review and import')).to_be_visible()
        check('import: badge column found under "Associate ID"', await page.locator('#col-employee_code').input_value() == '0')
        check('import preview: 3 new, 2 problems',
              await page.get_by_text('3 new').count() == 1 and await page.get_by_text('2 with problems (skipped)').count() == 1)
        await shot(page, '09b-import-preview')
        await page.get_by_role('button', name='Import 3 employees').click()
        await expect(page.get_by_role('heading', name='Import finished')).to_be_visible()
        await shot(page, '09c-import-done')
        check('import saved by badge with departments created',
              psql("select string_agg(employee_code || ' ' || last_name || ' ' || coalesce(d.name,'-'), ',' order by employee_code) from employees e left join departments d on d.id = e.department_id")
              == '10234 Lopez Picking,10235 Carter Receiving,10236 Nguyen Picking')
        check('import: hire date read from the Excel date cell', psql("select hire_date from employees where employee_code='10234'") == '2026-09-14')
        await page.get_by_role('link', name='See employees').click()

        # ---------- Employees: add one by hand, then edit ----------
        await page.get_by_role('link', name='Add employee').click()
        await page.get_by_label('Badge ID').fill('10240')
        await page.get_by_label('First name').fill('Rosa')
        await page.get_by_label('Last name').fill('Diaz')
        await page.get_by_label('Shift').select_option(label='First Shift')
        await page.get_by_label('Department').fill('Packing')
        await page.get_by_role('button', name='Add employee').click()
        await expect(page.get_by_role('heading', name='Employees')).to_be_visible()
        await expect(page.get_by_role('link', name='Diaz, Rosa')).to_be_visible()
        check('employee added with a new department', psql("select d.name from employees e join departments d on d.id=e.department_id where employee_code='10240'") == 'Packing')
        await page.get_by_role('link', name='Add employee').click()
        await page.get_by_label('Badge ID').fill('10234')
        await page.get_by_label('First name').fill('Copy'); await page.get_by_label('Last name').fill('Cat')
        await page.get_by_role('button', name='Add employee').click()
        await expect(page.get_by_role('alert')).to_contain_text('Another employee already has this badge ID')
        check('duplicate badge refused with a clear message', True)
        await page.get_by_role('link', name='Cancel').click()
        await page.get_by_label('Search').fill('carter')
        await expect(page.locator('.emp-table tbody tr')).to_have_count(1)
        check('employee search by name', True)
        await page.get_by_label('Search').fill('')
        await shot(page, '09d-employees')

        # ---------- Attendance ----------
        first = psql("select id from shifts where name='First Shift'")
        today = psql("select (now() at time zone 'America/New_York')::date")
        await page.goto(f'{BASE}/attendance?shift={first}&date={today}')
        await expect(page.get_by_role('heading', name='Attendance')).to_be_visible()
        await expect(page.locator('.roster-row')).to_have_count(3)
        check('attendance roster = active First Shift employees', True)
        await page.locator('#scan').fill('10234')
        await page.locator('#scan').press('Enter')
        await expect(page.locator('.scan-ok')).to_contain_text('Lopez, Maria')
        check('badge scan marks present or late',
              psql(f"select a.status from attendance_records a join employees e on e.id=a.employee_id where employee_code='10234' and work_date='{today}'") in ('present', 'late'))
        await page.locator('#scan').fill('99999'); await page.locator('#scan').press('Enter')
        await expect(page.locator('.scan-error')).to_contain_text('No active employee with badge 99999')
        check('unknown badge explained', True)
        carter = page.locator('.roster-row', has_text='Carter, James')
        await carter.get_by_role('button', name='Absent').click()
        await expect(carter.get_by_role('button', name='Absent')).to_have_attribute('aria-pressed', 'true')
        await carter.locator('select').select_option(label='No call / no show')
        await page.wait_for_timeout(600)
        check('absence saved with its reason',
              psql(f"select a.status || ' ' || reason_code from attendance_records a join employees e on e.id=a.employee_id where employee_code='10235' and work_date='{today}'") == 'absent no_call_no_show')
        await page.get_by_role('button', name=re.compile('Mark the 1 with no record')).click()
        await page.get_by_role('button', name='Yes, mark present').click()
        await expect(page.locator('.roster-row[data-status=none]')).to_have_count(0)
        check('mark the rest present', psql(f"select count(*) from attendance_records where work_date='{today}'") == '3')
        await shot(page, '09e-attendance')
        await page.set_viewport_size({'width': 390, 'height': 1000})
        await page.reload()
        await expect(page.locator('.roster-row')).to_have_count(3)
        await page.screenshot(path=f'{SHOTS}/09e2-mobile-attendance.png', scale='device')
        await page.set_viewport_size({'width': 1280, 'height': 860})
        diaz = page.locator('.roster-row', has_text='Diaz, Rosa')
        await diaz.get_by_role('button', name='Present').click()
        await expect(diaz).to_have_attribute('data-status', 'none')
        await page.wait_for_timeout(400)
        check('tapping the lit status clears it (No record)', psql(f"select count(*) from attendance_records where work_date='{today}'") == '2')
        await page.get_by_role('link', name='Today').click()
        await expect(page.get_by_text(re.compile(r'1 absent'))).to_be_visible()
        await shot(page, '09f-today-tally')
        check('Today shows the tally', True)

        # ---------- Reports ----------
        await page.get_by_role('link', name='Reports', exact=True).click()
        await page.locator('#r-period').select_option('today')
        await expect(page.locator('.kpi-main .kpi-value')).to_have_text('25%')
        check('report KPIs: 1 late of 4 scheduled = 25% attendance', await page.locator('.kpi-main .kpi-note').inner_text() == '1 of 4 scheduled')
        await shot(page, '09g-reports')
        await page.get_by_role('button', name=re.compile('^Absences')).click()
        await expect(page.locator('.report-table tbody tr')).to_have_count(1)
        check('absences report lists the absence with its reason', 'No call / no show' in await page.locator('.report-table').inner_text())
        async def grab(label):
            async with page.expect_download() as dl:
                await page.locator('.export').get_by_role('button', name=label).click()
            f = await dl.value
            path = os.path.join(SHOTS, f.suggested_filename); await f.save_as(path)
            return path
        csv_path = await grab('CSV')
        check('CSV export has the row', 'Carter, James' in open(csv_path, encoding='utf-8-sig').read() and csv_path.endswith('.csv'))
        xlsx_path = await grab('Excel')
        wbx = openpyxl.load_workbook(xlsx_path); vals = [c for row in wbx.active.iter_rows(values_only=True) for c in row if c]
        check('Excel export opens and has the row', 'Carter, James' in vals and 'Absences' in vals, str(vals[:12]))
        pdf_path = await grab('PDF')
        check('PDF export is a PDF', open(pdf_path, 'rb').read(5) == b'%PDF-')
        await page.get_by_role('button', name=re.compile('^No record')).click()
        await expect(page.locator('.report-table tbody tr')).to_have_count(2)
        check('no-record report lists the unmarked people', True)
        await page.get_by_role('button', name=re.compile('^By employee')).click()
        await shot(page, '09h-report-by-employee')

        # ---------- Account: change password, sign out, sign in ----------
        await page.get_by_role('link', name='My account').click()
        await page.get_by_label('New password', exact=True).fill(PW2)
        await page.get_by_label('Confirm new password').fill(PW2)
        await page.get_by_role('button', name='Change password').click()
        await expect(page.get_by_text('Password changed.')).to_be_visible()
        await shot(page, '10-account')
        await page.get_by_role('button', name='Sign out').click()
        await expect(page).to_have_url(re.compile(r'/login'))
        await page.get_by_label('Email').fill(BUYER)
        await page.get_by_label('Password').fill(PW1)
        await page.get_by_role('button', name='Sign in').click()
        await expect(page.get_by_role('alert')).to_contain_text('don’t match')
        check('old password stops working', True)
        await page.get_by_label('Password').fill(PW2)
        await page.get_by_role('button', name='Sign in').click()
        await expect(page.get_by_role('heading', name='Today', exact=True)).to_be_visible()
        check('new password signs in', True)

        # ---------- invite a supervisor from Settings → Team ----------
        await page.get_by_role('link', name='Settings', exact=True).click()
        await expect(page.get_by_role('heading', name='Team')).to_be_visible()
        await page.get_by_label('Their email').fill(SUPERVISOR)
        await page.get_by_role('button', name='Create invitation link').click()
        await expect(page.locator('.invite-link')).to_be_visible()
        link = await page.get_by_label('Invitation link').input_value()
        token = link.split('token=')[-1]
        check('owner can create an invitation link', len(token) == 64 and link.startswith(BASE + '/invite?token='), link)
        check('invitation is listed as waiting', await page.get_by_text('Waiting to accept').count() == 1)
        check('supervisor invite carries the warehouse', psql(f"select array_length(warehouse_ids,1) from invitations where email='{SUPERVISOR}'") == '1')
        await page.get_by_label('Their email').fill('mistake@example.com')
        await page.get_by_role('button', name='Create invitation link').click()
        await expect(page.locator('.team-list li', has_text='mistake@example.com')).to_be_visible()
        await page.locator('.team-list li', has_text='mistake@example.com').get_by_role('button', name='Cancel invitation').click()
        await expect(page.get_by_text('Invitation cancelled')).to_be_visible()
        check('cancelled invitation stops working', psql("select expires_at <= now() from invitations where email='mistake@example.com'") == 't')
        await shot(page, '10b-team')
        await page.get_by_role('button', name='Sign out').click()

        # ---------- forgot / reset password ----------
        await page.goto(f'{BASE}/forgot-password')
        await page.get_by_label('Email').fill(BUYER)
        t1 = time.time()
        await page.get_by_role('button', name='Send reset link').click()
        await expect(page.get_by_role('status')).to_contain_text('reset link is on its way')
        await shot(page, '11-forgot')
        await page.goto(mail_link(BUYER, 'recovery', t1))
        await expect(page.get_by_role('heading', name='Choose a new password')).to_be_visible(timeout=10000)
        await shot(page, '12-reset')
        await page.get_by_label('New password', exact=True).fill(PW3)
        await page.get_by_label('Confirm new password').fill(PW3)
        await page.get_by_role('button', name='Save new password').click()
        await expect(page.get_by_text('Password updated')).to_be_visible()
        await page.get_by_label('Email').fill(BUYER)
        await page.get_by_label('Password').fill(PW3)
        await page.get_by_role('button', name='Sign in').click()
        await expect(page.get_by_role('heading', name='Today', exact=True)).to_be_visible()
        check('reset link sets a new password', True)

        # ---------- two-step sign-in and the platform admin panel ----------
        await page.get_by_role('link', name='My account').click()
        await page.get_by_role('button', name='Set up two-step sign-in').click()
        secret = (await page.locator('.steps code').inner_text()).strip()
        await shot(page, '12b-two-step-setup')
        await page.get_by_label('Code').fill(totp(secret))
        await page.get_by_role('button', name='Turn on').click()
        await expect(page.get_by_text('Two-step sign-in is on.')).to_be_visible()
        check('two-step sign-in turned on', psql(f"select count(*) from auth.mfa_factors f join auth.users u on u.id=f.user_id where u.email='{BUYER}' and f.status='verified'") == '1')
        await page.get_by_role('button', name='Sign out').click()
        await page.get_by_label('Email').fill(BUYER)
        await page.get_by_label('Password').fill(PW3)
        await page.get_by_role('button', name='Sign in').click()
        await expect(page.get_by_role('heading', name='Enter your code')).to_be_visible()
        check('password alone is not enough once two-step is on', await page.get_by_role('heading', name='Today', exact=True).count() == 0)
        await shot(page, '12c-two-step-code')
        await page.get_by_label('Code').fill('000000')
        await page.get_by_role('button', name='Continue').click()
        await expect(page.get_by_role('alert')).to_contain_text('code didn’t work')
        check('wrong code refused', True)
        await page.wait_for_timeout(1000)
        await page.get_by_label('Code').fill(totp(secret))
        await page.get_by_role('button', name='Continue').click()
        await expect(page.get_by_role('heading', name='Today', exact=True)).to_be_visible()
        check('right code signs in', True)

        psql(f"insert into platform_admins (user_id) select id from auth.users where email='{BUYER}'")
        await page.goto(f'{BASE}/admin')
        await expect(page.get_by_role('heading', name='Admin', exact=True)).to_be_visible()
        await expect(page.locator('.report-table').first).to_contain_text('Northgate Logistics LLC')
        check('platform admin sees companies with two-step sign-in', True)
        check('admin panel shows the Hotmart deliveries', await page.get_by_text('PURCHASE_APPROVED').count() >= 1)
        await page.get_by_role('button', name='Give a license').click()
        await page.get_by_label('Buyer email').fill('pilot@example.com')
        await page.get_by_role('button', name='Save license').click()
        await expect(page.get_by_text('License created.')).to_be_visible()
        check('admin gives a pilot license', psql("select status || ' ' || plan_code from licenses where buyer_email='pilot@example.com'") == 'trial professional')
        await shot(page, '12d-admin')
        await page.goto(f'{BASE}/account')
        await page.get_by_role('button', name='Turn off').click()
        await page.get_by_role('button', name='Yes, turn off').click()
        await expect(page.get_by_text('Two-step sign-in is off.')).to_be_visible()
        await page.get_by_role('button', name='Sign out').click()

        # ---------- someone without a purchase ----------
        ctx2 = await browser.new_context(viewport={'width': 390, 'height': 844})
        p2 = await ctx2.new_page()
        OPEN_PAGES.append(p2)
        await p2.goto(f'{BASE}/signup')
        await p2.get_by_label('Your name').fill('No Purchase')
        await p2.get_by_label('Email').fill(STRANGER)
        await p2.get_by_label('Password', exact=True).fill(PW1)
        await p2.get_by_label('Confirm password').fill(PW1)
        t2 = time.time()
        await p2.get_by_role('button', name='Create account').click()
        await expect(p2.get_by_role('heading', name='Check your email')).to_be_visible()
        await p2.goto(mail_link(STRANGER, 'signup', t2))
        await p2.get_by_label('Company name').fill('Freeloader Inc')
        await p2.get_by_role('button', name='Continue').click()
        await expect(p2.get_by_role('heading', name=re.compile('can’t find a purchase'))).to_be_visible()
        await shot(p2, '13-mobile-no-purchase')
        check('no purchase: company is not created', psql("select count(*) from organizations where name='Freeloader Inc'") == '0')

        # ---------- supervisor accepts the invitation ----------
        ctx3 = await browser.new_context(viewport={'width': 390, 'height': 844})
        p3 = await ctx3.new_page()
        OPEN_PAGES.append(p3)
        await p3.goto(f'{BASE}/invite?token={token}')
        await expect(p3.get_by_role('heading', name='You’ve been invited')).to_be_visible()
        await shot(p3, '14-mobile-invite')
        await p3.get_by_role('link', name='Create an account').click()
        await p3.get_by_label('Your name').fill('Sam Supervisor')
        await p3.get_by_label('Email').fill(SUPERVISOR)
        await p3.get_by_label('Password', exact=True).fill(PW1)
        await p3.get_by_label('Confirm password').fill(PW1)
        t3 = time.time()
        await p3.get_by_role('button', name='Create account').click()
        await expect(p3.get_by_role('heading', name='Check your email')).to_be_visible()
        await p3.goto(mail_link(SUPERVISOR, 'signup', t3))
        await expect(p3.get_by_role('heading', name='Today', exact=True)).to_be_visible(timeout=10000)
        await shot(p3, '15-mobile-supervisor-today')
        check('invitation joins the supervisor to the company',
              psql(f"select m.role from memberships m join auth.users u on u.id=m.user_id where u.email='{SUPERVISOR}'") == 'supervisor')
        check('supervisor sees no Settings link', await p3.get_by_role('link', name='Settings').count() == 0)
        await p3.get_by_role('link', name='Attendance', exact=True).click()
        await expect(p3.locator('.roster-row').first).to_be_visible()
        await shot(p3, '15b-mobile-attendance')
        check('supervisor can take attendance in her warehouse', await p3.locator('.roster-row').count() >= 1)
        await p3.get_by_role('link', name='Employees', exact=True).click()
        await expect(p3.get_by_role('heading', name='Employees')).to_be_visible()
        check('supervisor sees employees read-only', await p3.get_by_role('link', name='Add employee').count() == 0)
        await shot(p3, '15c-mobile-employees')
        await p3.get_by_role('link', name='Reports', exact=True).click()
        await expect(p3.locator('.kpi-main .kpi-value')).to_be_visible()
        await shot(p3, '15d-mobile-reports')
        await p3.goto(f'{BASE}/settings')
        await expect(p3).to_have_url(re.compile(r'/$'))
        check('supervisor is sent away from Settings', True)

        # ---------- refund locks the company ----------
        out = psql("select billing.apply_hotmart_event('e2e-2','PURCHASE_REFUNDED', now(), 'SUB-E2E', null, "
                   f"'{BUYER}', null, null)")
        check('refund webhook cancels the license', out == 'license_cancelled', out)
        await p3.reload()
        await expect(p3.get_by_role('heading', name='Your subscription is inactive')).to_be_visible()
        await shot(p3, '16-mobile-inactive')
        await page.goto(f'{BASE}/login')
        await page.get_by_label('Email').fill(BUYER)
        await page.get_by_label('Password').fill(PW3)
        await page.get_by_role('button', name='Sign in').click()
        await expect(page.get_by_role('heading', name='Your subscription is inactive')).to_be_visible()
        await expect(page.get_by_text(re.compile('30 days after the subscription ends'))).to_be_visible()
        check('owner sees the export-window note', True)
        await shot(page, '17-owner-inactive')


        # ---------- live demo for visitors ----------
        ctx4 = await browser.new_context(viewport={'width': 390, 'height': 844})
        p4 = await ctx4.new_page()
        OPEN_PAGES.append(p4)
        await p4.goto(f'{BASE}/login')
        await p4.get_by_role('link', name='Open the live demo').click()
        await expect(p4.get_by_role('heading', name='Try it with sample data')).to_be_visible()
        await shot(p4, '18-mobile-demo-start')
        await p4.get_by_role('button', name='Open the demo').click()
        await expect(p4.get_by_role('heading', name='Today', exact=True)).to_be_visible(timeout=15000)
        await expect(p4.locator('.demo-bar')).to_contain_text('Demo with sample data')
        check('demo opens a sandbox company', await p4.get_by_text('Demo Logistics').count() >= 1)
        await shot(p4, '19-mobile-demo-today')
        await p4.get_by_role('link', name='Reports', exact=True).click()
        await p4.locator('#r-period').select_option('last_week')
        await expect(p4.locator('.kpi-main .kpi-value')).not_to_have_text('—')
        check('demo reports show sample attendance', int((await p4.locator('.kpi-note').inner_text()).split(' of ')[1].split()[0]) > 100)
        await shot(p4, '20-mobile-demo-reports')
        check('demo company is marked as demo and kept apart', psql("select count(*) from organizations where is_demo") == '1')
        await p4.goto(f'{BASE}/account')
        await expect(p4.get_by_role('heading', name='You’re using the demo')).to_be_visible()
        await p4.locator('.panel').get_by_role('button', name='Leave demo').click()
        await expect(p4).to_have_url(re.compile(r'/login'))
        check('visitor can leave the demo', True)

        # ---------- legal pages ----------
        await p4.get_by_role('link', name='Terms').click()
        await expect(p4.get_by_role('heading', name='Terms of Service')).to_be_visible()
        await p4.get_by_role('link', name='Privacy').first.click()
        await expect(p4.get_by_role('heading', name='Privacy Policy')).to_be_visible()
        await shot(p4, '21-mobile-privacy')
        check('terms and privacy pages open', True)

        check('no JavaScript errors', not console_errors, '; '.join(console_errors))

async def run():
    PW.append(await async_playwright().start())
    try:
        await main()
    except Exception:
        for i, pg in enumerate(OPEN_PAGES):
            try:
                await pg.screenshot(path=f'{SHOTS}/FAIL-{i}.png', full_page=True)
                print(f'FAIL-{i} at', pg.url)
            except Exception:
                pass
        raise
    finally:
        await PW[0].stop()

asyncio.run(run())
print(f'\n{len(passed)} passed, {len(failed)} failed')
for f in failed:
    print(' ✗', f)
sys.exit(1 if failed else 0)
